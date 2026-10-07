"""ExcelService testleri: örnek workbook oluşturma ve toplu dışa aktarım."""
from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.models.category import Category
from app.models.formula import Difficulty, Formula
from app.services.excel_service import ExcelService, _safe_sheet_name, to_universal_formula


# ---- Fixtures ----
@pytest.fixture()
def excel_service(tmp_path: Path) -> ExcelService:
    return ExcelService(export_dir=tmp_path)


@pytest.fixture()
def second_formula(session, sample_category) -> Formula:
    formula = Formula(
        category_id=sample_category.id,
        name_tr="EĞER",
        name_en="IF",
        syntax_tr="=EĞER(test;_doğru;yanlış)",
        short_description_tr="Bir koşulu test eder.",
        detailed_explanation_tr="EĞER, verilen koşulun DOĞRU veya YANLIŞ olduğunu kontrol eder.",
        example_formula='=EĞER(A1>10;"büyük";"küçük")',
        example_description_tr="A1 10'dan büyükse 'büyük', değilse 'küçük' yazar.",
        example_result="büyük",
        difficulty=Difficulty.ORTA,
    )
    session.add(formula)
    session.commit()
    return formula


# ---- _safe_sheet_name Tests ----
class TestSafeSheetName:
    def test_normal_name_unchanged(self):
        assert _safe_sheet_name("Matematik") == "Matematik"

    def test_truncates_to_31_chars(self):
        long_name = "A" * 40
        result = _safe_sheet_name(long_name)
        assert len(result) <= 31

    def test_removes_forbidden_chars(self):
        result = _safe_sheet_name('Test[0]:*?/\\Name')
        assert not any(ch in result for ch in '[]:*?/\\')

    def test_empty_returns_default(self):
        assert _safe_sheet_name("") == "Sayfa1"
        assert _safe_sheet_name("[]:*?/\\") == "Sayfa1"


# ---- build_example_workbook Tests ----
class TestBuildExampleWorkbook:
    def test_creates_xlsx_file(self, excel_service, sample_formula):
        path = excel_service.build_example_workbook(sample_formula)
        assert path.exists()
        assert path.suffix == ".xlsx"

    def test_filename_contains_formula_name(self, excel_service, sample_formula):
        path = excel_service.build_example_workbook(sample_formula)
        assert "TOPLA" in path.name

    def test_workbook_has_correct_sheet_title(self, excel_service, sample_formula):
        path = excel_service.build_example_workbook(sample_formula)
        wb = load_workbook(path)
        assert wb.sheetnames[0] == "TOPLA"

    def test_workbook_contains_formula_name_in_cell(self, excel_service, sample_formula):
        path = excel_service.build_example_workbook(sample_formula)
        wb = load_workbook(path)
        ws = wb.active
        # Başlık, veri bloğunun sağındaki bilgi bloğununundur (A1 örnek veriye ait).
        assert "TOPLA" in str(ws["C1"].value)
        assert "SUM" in str(ws["C1"].value)

    def test_workbook_contains_syntax(self, excel_service, sample_formula):
        path = excel_service.build_example_workbook(sample_formula)
        wb = load_workbook(path)
        ws = wb.active
        # Bilgi bloğu: C sütunu etiketler, D sütunu değerler.
        assert ws["C3"].value == "Söz Dizimi:"
        assert ws["D3"].value == sample_formula.syntax_tr

    def test_workbook_contains_sample_data(self, excel_service, sample_formula):
        path = excel_service.build_example_workbook(sample_formula)
        wb = load_workbook(path)
        ws = wb.active
        # Örnek veriler, formülün referans ettiği A1:A5 aralığına yazılır.
        assert ws["A1"].value == 12
        assert ws["A5"].value == 8
        assert [ws[f"A{row}"].value for row in range(1, 6)] == [12, 23, 35, 10, 8]

    def test_workbook_contains_example_formula(self, excel_service, sample_formula):
        path = excel_service.build_example_workbook(sample_formula)
        wb = load_workbook(path)
        ws = wb.active
        # Sonuç satırı, veri bloğunun hemen altındadır (A1:A5 -> 5. satır + 2).
        result_row = 5 + 2
        assert ws[f"A{result_row}"].value == "Sonuç →"
        # Formül dosya biçiminin beklediği EVRENSEL gramerle saklanır
        # (=TOPLA -> =SUM); Türkçe Excel açınca =TOPLA olarak gösterir.
        assert ws[f"B{result_row}"].value == "=SUM(A1:A5)"

    def test_workbook_shows_expected_result_and_metadata(self, excel_service, sample_formula):
        path = excel_service.build_example_workbook(sample_formula)
        wb = load_workbook(path)
        ws = wb.active
        assert ws["C11"].value == "Beklenen Sonuç:"
        assert ws["D11"].value == "88"
        assert ws["D4"].value == "Matematik"          # kategori
        assert ws["D5"].value == "Başlangıç"          # zorluk
        # Formülün baktığı aralık ile bilgi bloğu çakışmamalıdır:
        assert ws["A3"].value == 35                     # veri, metin değil

    def test_sample_data_written_to_all_referenced_ranges(self, excel_service, sample_category, session):
        formula = Formula(
            category_id=sample_category.id,
            name_tr="ÇOKETOPLA",
            name_en="SUMIFS",
            syntax_tr="=ÇOKETOPLA(...) ",
            short_description_tr="Çoklu koşul toplama.",
            detailed_explanation_tr="Açıklama.",
            example_formula="=ÇOKETOPLA(C1:C5;A1:A5;B1:B5)",
            example_description_tr="Örnek.",
            difficulty=Difficulty.ORTA,
        )
        session.add(formula)
        session.commit()
        path = excel_service.build_example_workbook(formula)
        wb = load_workbook(path)
        ws = wb.active
        # A ve C sütunları havuz değerleriyle dolu:
        for col in ("A", "C"):
            assert ws[f"{col}1"].value == 12
            assert ws[f"{col}5"].value == 8
        # B sütunu özel: ilk hücre arama anahtarıdır (A1 ile aynı),
        # sonrakiler +60 ile koşullu kriterlerin (\">50\") eşleşmesini sağlar.
        assert ws["B1"].value == 12
        assert ws["B5"].value == 68
        # Bilgi bloğu veri bloğunun sağına kayar (E sütunundan başlar).
        assert ws["E1"].value is not None
        assert "ÇOKETOPLA" in str(ws["E1"].value)

    def test_formula_without_cell_references(self, excel_service, sample_category, session):
        formula = Formula(
            category_id=sample_category.id,
            name_tr="YUVARLA",
            name_en="ROUND",
            syntax_tr="=YUVARLA(sayı; basamak)",
            short_description_tr="Sayıyı yuvarlar.",
            detailed_explanation_tr="Açıklama.",
            example_formula="=YUVARLA(17,6456;2)",
            example_description_tr="17,6456 sayısını 2 basamağa yuvarlar.",
            example_result="17,65",
            difficulty=Difficulty.BASLANGIC,
        )
        session.add(formula)
        session.commit()
        path = excel_service.build_example_workbook(formula)
        wb = load_workbook(path)
        ws = wb.active
        values = [cell.value for row in ws.iter_rows() for cell in row]
        formulas = [
            cell.value for row in ws.iter_rows() for cell in row if cell.data_type == "f"
        ]
        assert "YUVARLA" in str(ws["A1"].value)       # başlık A1'de
        # Canlı formül evrensel gramerle saklanır:
        assert formulas == ["=ROUND(17.6456,2)"]
        assert "17,65" in values                        # beklenen sonuç gösterilir

    def test_time_result_gets_time_number_format(self, excel_service, sample_category, session):
        """Saat/tarih sonucu ham ondalık sayı yerine doğru biçimde görünmeli."""
        formula = Formula(
            category_id=sample_category.id,
            name_tr="ZAMAN",
            name_en="TIME",
            syntax_tr="=ZAMAN(saat; dakika; saniye)",
            short_description_tr="Saat üretir.",
            detailed_explanation_tr="Açıklama.",
            example_formula="=ZAMAN(14;30;0)",
            example_description_tr="14:30 saatindeki zaman değerini üretir.",
            example_result="14:30:00",
            difficulty=Difficulty.BASLANGIC,
        )
        session.add(formula)
        session.commit()
        path = excel_service.build_example_workbook(formula)
        ws = load_workbook(path).active
        assert ws["B17"].number_format == "hh:mm:ss"
        assert ws["B17"].data_type == "f"

    def test_filename_is_sanitized(self, excel_service, sample_category, session):
        formula = Formula(
            category_id=sample_category.id,
            name_tr='TE/ST<FORMÜL',
            name_en="TEST",
            syntax_tr="=TEST()",
            short_description_tr="Test.",
            detailed_explanation_tr="Açıklama.",
            example_formula="=TEST(A1:A5)",
            example_description_tr="Örnek.",
            difficulty=Difficulty.BASLANGIC,
        )
        session.add(formula)
        session.commit()
        path = excel_service.build_example_workbook(formula)
        assert path.exists()
        assert not any(ch in path.name for ch in '<>:"/\\|?*')
        # Sekme adı da Excel kurallarına uymalı:
        wb = load_workbook(path)
        assert not any(ch in wb.sheetnames[0] for ch in '[]:*?/\\')

    def test_syntax_stored_as_text_not_formula(self, excel_service, sample_formula):
        """Söz dizimi metni ('=TOPLA(sayı1; [sayı2]; ...)') formül olarak
        saklanırsa Excel dosyanın AÇILIŞINDA tüm kitabı reddeder.
        Metin olarak yazıldığı doğrulanır."""
        path = excel_service.build_example_workbook(sample_formula)
        ws = load_workbook(path).active
        assert ws["D3"].data_type != "f"          # formül değil
        assert ws["D3"].value == sample_formula.syntax_tr

    def test_live_result_cell_remains_formula(self, excel_service, sample_formula):
        """Canlı örnek hücresi (B7) gerçek formül olarak kalmalıdır —
        Excel'de açıldığında 88 hesaplar."""
        path = excel_service.build_example_workbook(sample_formula)
        ws = load_workbook(path).active
        assert ws["B7"].data_type == "f"

    def test_live_formula_uses_universal_grammar(self, excel_service, sample_category, session):
        """OOXML biçimi Türkçe formül gramerini KABUL ETMEZ:
        ';' ayracı dosyanın tamamını reddettirir, Türkçe fonksiyon adı
        #NAME? verir. Canlı hücreye kanonik biçim yazılır."""
        cases = [
            ("EĞER", "IF", '=EĞER(A1>10;"küçük";"büyük")', '=IF(A1>10,"küçük","büyük")'),
            ("YUVARLA", "ROUND", "=YUVARLA(17,6456;2)", "=ROUND(17.6456,2)"),
            ("DOLAYLI", "INDIRECT", '=DOLAYLI("A1")', '=INDIRECT("A1")'),
            ("ETOPLA", "SUMIF", '=ETOPLA(A1:A5;\">20\")', '=SUMIF(A1:A5,\">20\")'),
        ]
        for name_tr, name_en, example, expected in cases:
            formula = Formula(
                category_id=sample_category.id,
                name_tr=name_tr,
                name_en=name_en,
                syntax_tr=f"={name_tr}(...)",
                short_description_tr="Test.",
                detailed_explanation_tr="Açıklama.",
                example_formula=example,
                example_description_tr="Örnek.",
                difficulty=Difficulty.BASLANGIC,
            )
            session.add(formula)
            session.commit()
            path = excel_service.build_example_workbook(formula)
            ws = load_workbook(path).active
            formulas_in_sheet = [
                c.value for row in ws.iter_rows() for c in row if c.data_type == "f"
            ]
            assert formulas_in_sheet == [expected], f"{name_tr}: {formulas_in_sheet}"
            # Tırnak içi metin (ör. \">20\") korunmuş olmalı:
            assert expected in str(formulas_in_sheet[0])

    def test_placeholder_example_formula_written_as_text(self, excel_service, sample_category, session):
        """Yer tutuculu ('...') formül geçersiz sözdizimidir; Excel dosyayı
        reddeder. Bu yüzden düz metin olarak yazılır."""
        formula = Formula(
            category_id=sample_category.id,
            name_tr="DENEME",
            name_en="DEMO",
            syntax_tr="=DENEME(...)",
            short_description_tr="Test.",
            detailed_explanation_tr="Açıklama.",
            example_formula="=DENEME(...)",
            example_description_tr="Örnek.",
            difficulty=Difficulty.BASLANGIC,
        )
        session.add(formula)
        session.commit()
        path = excel_service.build_example_workbook(formula)
        ws = load_workbook(path).active
        values = [c for row in ws.iter_rows() for c in row if c.value == "=DENEME(...)"]
        assert values, "yer tutucu formül sayfada bulunamadı"
        assert all(c.data_type != "f" for c in values), "yer tutucu formül olarak saklanmış"

    def test_turkish_name_with_special_chars(self, excel_service, sample_category, session):
        formula = Formula(
            category_id=sample_category.id,
            name_tr="ÇOKETOPLA",
            name_en="SUMIFS",
            syntax_tr="=ÇOKETOPLA(...)",
            short_description_tr="Test.",
            detailed_explanation_tr="Açıklama.",
            example_formula="=ÇOKETOPLA(A1:A5)",
            example_description_tr="Örnek.",
            difficulty=Difficulty.ORTA,
        )
        session.add(formula)
        session.commit()
        path = excel_service.build_example_workbook(formula)
        assert path.exists()
        assert "ÇOKETOPLA" in path.name


# ---- to_universal_formula (dosya biçimi çevirisi) Tests ----
class TestToUniversalFormula:
    def test_function_names_become_canonical(self):
        assert to_universal_formula("=TOPLA(A1:A5)") == "=SUM(A1:A5)"
        assert to_universal_formula("=YUVARLA(17,6456;2)") == "=ROUND(17.6456,2)"

    def test_boolean_literals(self):
        assert to_universal_formula('=EĞER(DOĞRU;1;2)') == "=IF(TRUE,1,2)"
        assert to_universal_formula('=DÜŞEYARA(A1;B1:D10;3;YANLIŞ)') == (
            "=VLOOKUP(A1,B1:D10,3,FALSE)"
        )

    def test_post_2007_functions_get_xlfn_prefix(self):
        """2007 sonrası fonksiyonlar dosya biçiminde '_xlfn.' öneki ister;
        öneksiz ad Excel'de #NAME? üretir (dosya açılır ama sonuç hatalı)."""
        assert to_universal_formula('=ÇOKEĞER(A1>=50;"CB";DOĞRU;"FF")') == (
            '=_xlfn.IFS(A1>=50,"CB",TRUE,"FF")'
        )
        assert to_universal_formula('=EĞERYOKSA(YOKSAY();"yok")') == (
            '=_xlfn.IFNA(NA(),"yok")'
        )
        assert to_universal_formula('=METİNBİRLEŞTİR(", ";DOĞRU;A1:A5)') == (
            '=_xlfn.TEXTJOIN(", ",TRUE,A1:A5)'
        )
        assert to_universal_formula('=ÇOKEĞERMAK(C1:C5;A1:A5;">10")') == (
            '=_xlfn.MAXIFS(C1:C5,A1:A5,">10")'
        )
        # 2007 ve öncesi önek istemez:
        assert to_universal_formula("=TOPLA(A1:A5)") == "=SUM(A1:A5)"
        assert to_universal_formula('=ETOPLA(A1:A5;">20")') == '=SUMIF(A1:A5,">20")'

    def test_array_constants_keep_comma_separator(self):
        # Dizi sabitindeki virgül AYIRAÇTIR; ondalığa çevrilemez!
        assert to_universal_formula("=İÇ_VERİM_ORANI({-10000;3000;4000})") == (
            "=IRR({-10000,3000,4000})"
        )

    def test_string_literals_untouched(self):
        assert to_universal_formula('=BUL("a;b,1";A1)') == '=FIND("a;b,1",A1)'
        assert to_universal_formula('=EĞER(A1;"DOĞRU";"YANLIŞ")') == (
            '=IF(A1,"DOĞRU","YANLIŞ")'   # tırnak içindeki metin değişmez
        )

    def test_own_names_override(self):
        assert to_universal_formula("=ÖZELFORMÜL(A1)", {"ÖZELFORMÜL": "MYFUNC"}) == (
            "=MYFUNC(A1)"
        )


# ---- export_all_formulas Tests ----
class TestExportAllFormulas:
    def test_creates_reference_workbook(self, excel_service, sample_formula):
        path = excel_service.export_all_formulas([sample_formula])
        assert path.exists()
        assert path.suffix == ".xlsx"
        assert "referans" in path.name

    def test_workbook_has_category_sheet(self, excel_service, sample_formula):
        path = excel_service.export_all_formulas([sample_formula])
        wb = load_workbook(path)
        # Should have a sheet named after the category
        assert len(wb.sheetnames) >= 1

    def test_workbook_contains_formula_data(self, excel_service, sample_formula):
        path = excel_service.export_all_formulas([sample_formula])
        wb = load_workbook(path)
        ws = wb.active
        # Row 2 should have formula data (row 1 is header)
        assert ws.cell(row=2, column=1).value == "TOPLA"
        assert ws.cell(row=2, column=2).value == "SUM"

    def test_multiple_categories_separate_sheets(
        self, excel_service, sample_formula, second_formula
    ):
        path = excel_service.export_all_formulas([sample_formula, second_formula])
        wb = load_workbook(path)
        # Both formulas are in the same category, so 1 sheet
        assert len(wb.sheetnames) == 1

    def test_turkish_name_in_category_sheet(self, excel_service, sample_category, session):
        formula = Formula(
            category_id=sample_category.id,
            name_tr="ÇOKETOPLA",
            name_en="SUMIFS",
            syntax_tr="=ÇOKETOPLA(...)",
            short_description_tr="Çoklu koşul toplama.",
            detailed_explanation_tr="Açıklama.",
            example_formula="=ÇOKETOPLA(...)",
            example_description_tr="Örnek.",
            difficulty=Difficulty.ORTA,
        )
        session.add(formula)
        session.commit()
        path = excel_service.export_all_formulas([formula])
        wb = load_workbook(path)
        ws = wb.active
        # Turkish name should be in column A, row 2
        assert ws.cell(row=2, column=1).value == "ÇOKETOPLA"

    def test_headers_are_present(self, excel_service, sample_formula):
        path = excel_service.export_all_formulas([sample_formula])
        wb = load_workbook(path)
        ws = wb.active
        expected_headers = ["Türkçe Ad", "İngilizce Ad", "Söz Dizimi", "Açıklama", "Örnek", "Zorluk"]
        for col_idx, header in enumerate(expected_headers, start=1):
            assert ws.cell(row=1, column=col_idx).value == header

    def test_rich_headers_and_values_present(self, excel_service, sample_formula):
        path = excel_service.export_all_formulas([sample_formula])
        wb = load_workbook(path)
        ws = wb.active
        # Zengin sütunlar (6'dan sonra) başlık ve değerleriyle yerinde olmalı:
        assert ws.cell(row=1, column=7).value == "Detaylı Açıklama"
        assert ws.cell(row=1, column=9).value == "Beklenen Sonuç"
        assert ws.cell(row=1, column=13).value == "Etiketler"
        assert ws.cell(row=2, column=7).value == sample_formula.detailed_explanation_tr
        assert ws.cell(row=2, column=9).value == "88"
        assert ws.cell(row=2, column=13).value == "matematik, toplama"

    def test_autofilter_and_freeze_panes(self, excel_service, sample_formula):
        path = excel_service.export_all_formulas([sample_formula])
        wb = load_workbook(path)
        ws = wb.active
        assert ws.freeze_panes == "A2"
        assert ws.auto_filter.ref is not None
        assert ws.auto_filter.ref.startswith("A1:")

    def test_empty_list_creates_valid_workbook(self, excel_service):
        path = excel_service.export_all_formulas([])
        assert path.exists()
        assert path.suffix == ".xlsx"
        wb = load_workbook(path)
        # Sıfır sayfalık workbook Excel'de açılamaz; en az bir sayfa zorunlu:
        assert len(wb.sheetnames) >= 1
        assert wb.active["A1"].value is not None

    def test_formula_like_columns_are_text(self, excel_service, sample_formula):
        """Söz Dizimi (3), Örnek (5) ve Alıştırma Cevabı (11) sütunları
        `=` ile başlar ve formül değil METİN olarak saklanmalıdır:
        aksi hâlde Excel dosyayı reddeder veya satırlar rastgele hesaplanır."""
        path = excel_service.export_all_formulas([sample_formula])
        ws = load_workbook(path).active
        for col in (3, 5, 11):
            cell = ws.cell(row=2, column=col)
            assert cell.data_type != "f", f"sütun {col} formül olarak saklanmış"
            assert str(cell.value).startswith("=")
        # Türkçe Ad sütunu gibi normal metin sütunları da etkilenmemeli:
        assert ws.cell(row=2, column=1).value == "TOPLA"
